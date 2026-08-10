import io
from uuid import UUID
from datetime import datetime
from typing import Dict, Any, List
from sqlalchemy.orm import Session
from sqlalchemy import and_

from app.models import HabitScore, UserBehaviorAnalytic, ChallengeResult, User
from app.utils import parse_time_to_minutes

# openpyxl for Excel generation
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# reportlab for PDF generation
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

def generate_report_data(
    user_id: UUID,
    report_type: str,
    start_date: str,  # YYYY-MM-DD
    end_date: str,    # YYYY-MM-DD
    db: Session
) -> Dict[str, Any]:
    """Compiles statistics and list data for reports."""
    report_type = report_type.lower()
    
    # 1. Fetch Habit Scores
    scores_query = db.query(HabitScore).filter(
        HabitScore.user_id == user_id,
        HabitScore.date >= start_date,
        HabitScore.date <= end_date
    ).order_by(HabitScore.date.asc())
    scores = scores_query.all()

    # 2. Fetch Behavior Analytics
    behavior_query = db.query(UserBehaviorAnalytic).filter(
        UserBehaviorAnalytic.user_id == user_id,
        UserBehaviorAnalytic.date >= start_date,
        UserBehaviorAnalytic.date <= end_date
    ).order_by(UserBehaviorAnalytic.date.asc())
    behavior_logs = behavior_query.all()

    # 3. Fetch Challenge Results
    challenges_query = db.query(ChallengeResult).filter(
        ChallengeResult.user_id == user_id,
        ChallengeResult.solved_at >= datetime.strptime(start_date, "%Y-%m-%d"),
        ChallengeResult.solved_at <= datetime.strptime(end_date + " 23:59:59", "%Y-%m-%d %H:%M:%S")
    ).order_by(ChallengeResult.solved_at.asc())
    challenge_results = challenges_query.all()

    if report_type == "habit":
        total_days = len(scores)
        if total_days > 0:
            avg_consistency = sum(s.wake_up_consistency for s in scores) / total_days
            avg_completion = sum(s.challenge_completion for s in scores) / total_days
            avg_snooze = sum(s.snooze_reduction for s in scores) / total_days
            avg_sleep = sum(s.sleep_adherence for s in scores) / total_days
            avg_overall = sum(s.overall_score for s in scores) / total_days
        else:
            avg_consistency = avg_completion = avg_snooze = avg_sleep = avg_overall = 0.0
            
        return {
            "report_name": "Habit Performance Report",
            "start_date": start_date,
            "end_date": end_date,
            "summary": {
                "total_tracked_days": total_days,
                "average_wake_up_consistency": round(avg_consistency, 2),
                "average_challenge_completion": round(avg_completion, 2),
                "average_snooze_reduction": round(avg_snooze, 2),
                "average_sleep_adherence": round(avg_sleep, 2),
                "average_overall_score": round(avg_overall, 2),
            },
            "data": [
                {
                    "date": s.date,
                    "wake_up_consistency": s.wake_up_consistency,
                    "challenge_completion": s.challenge_completion,
                    "snooze_reduction": s.snooze_reduction,
                    "sleep_adherence": s.sleep_adherence,
                    "overall_score": s.overall_score
                } for s in scores
            ]
        }

    elif report_type == "wake-up":
        total_logs = len(behavior_logs)
        total_snoozes = sum(log.snooze_count for log in behavior_logs)
        avg_delay = sum(log.wake_up_delay for log in behavior_logs) / total_logs if total_logs > 0 else 0.0
        solved_challenges = sum(1 for log in behavior_logs if log.challenge_solved)
        
        return {
            "report_name": "Wake-up Patterns & Consistency Report",
            "start_date": start_date,
            "end_date": end_date,
            "summary": {
                "total_wakeups_tracked": total_logs,
                "total_snooze_count": total_snoozes,
                "average_snooze_count": round(total_snoozes / total_logs, 2) if total_logs > 0 else 0.0,
                "average_wake_up_delay_seconds": round(avg_delay, 1),
                "challenge_solve_rate_percentage": round((solved_challenges / total_logs) * 100, 1) if total_logs > 0 else 0.0
            },
            "data": [
                {
                    "date": log.date,
                    "target_wake_up_time": log.target_wake_up_time,
                    "wake_up_time": log.wake_up_time,
                    "wake_up_delay_seconds": log.wake_up_delay,
                    "snooze_count": log.snooze_count,
                    "challenge_solved": log.challenge_solved,
                    "challenge_solve_time_seconds": log.challenge_solve_time
                } for log in behavior_logs
            ]
        }

    elif report_type == "challenge":
        total_solved = len(challenge_results)
        avg_time = sum(r.completion_time for r in challenge_results) / total_solved if total_solved > 0 else 0.0
        avg_accuracy = sum(r.accuracy for r in challenge_results) / total_solved if total_solved > 0 else 0.0
        
        category_breakdown = {}
        for r in challenge_results:
            cat_name = r.challenge.category.name if (r.challenge and r.challenge.category) else "Other"
            category_breakdown[cat_name] = category_breakdown.get(cat_name, 0) + 1
            
        return {
            "report_name": "Cognitive Challenge Performance Report",
            "start_date": start_date,
            "end_date": end_date,
            "summary": {
                "total_challenges_solved": total_solved,
                "average_completion_time_seconds": round(avg_time, 1),
                "average_accuracy_percentage": round(avg_accuracy * 100, 1),
                "category_breakdown": category_breakdown
            },
            "data": [
                {
                    "solved_at": r.solved_at.isoformat(),
                    "category": r.challenge.category.name if (r.challenge and r.challenge.category) else "Other",
                    "difficulty": r.challenge.difficulty,
                    "score": r.score,
                    "accuracy_percentage": round(r.accuracy * 100, 1),
                    "completion_time_seconds": r.completion_time,
                    "attempts": r.total_attempts
                } for r in challenge_results
            ]
        }

    elif report_type == "productivity":
        total_days = len(scores)
        high_productive = sum(1 for s in scores if s.overall_score >= 80 and s.snooze_reduction >= 85)
        med_productive = sum(1 for s in scores if 60 <= s.overall_score < 80)
        low_productive = sum(1 for s in scores if s.overall_score < 60)
        
        efficiency_index = 0.0
        if total_days > 0:
            # Formula: weighted average of overall score and snooze reduction
            avg_overall = sum(s.overall_score for s in scores) / total_days
            avg_snooze_red = sum(s.snooze_reduction for s in scores) / total_days
            efficiency_index = 0.7 * avg_overall + 0.3 * avg_snooze_red
            
        return {
            "report_name": "Productivity & Morning Routine Efficiency Report",
            "start_date": start_date,
            "end_date": end_date,
            "summary": {
                "total_days_evaluated": total_days,
                "morning_efficiency_index": round(efficiency_index, 2),
                "high_productivity_days": high_productive,
                "medium_productivity_days": med_productive,
                "low_productivity_days": low_productive,
                "high_productivity_ratio_percentage": round((high_productive / total_days) * 100, 1) if total_days > 0 else 0.0
            },
            "data": [
                {
                    "date": s.date,
                    "overall_score": s.overall_score,
                    "snooze_reduction": s.snooze_reduction,
                    "wake_up_consistency": s.wake_up_consistency,
                    "routine_status": "High Productivity" if s.overall_score >= 80 and s.snooze_reduction >= 85 else ("Medium Productivity" if 60 <= s.overall_score < 80 else "Low Productivity")
                } for s in scores
            ]
        }

    elif report_type == "sleep":
        total_logs = len(behavior_logs)
        avg_sleep = sum(log.sleep_duration for log in behavior_logs if log.sleep_duration is not None) / total_logs if total_logs > 0 else 0.0
        adherence_days = sum(1 for s in scores if s.sleep_adherence >= 80)
        
        return {
            "report_name": "Sleep Schedule Analytics Report",
            "start_date": start_date,
            "end_date": end_date,
            "summary": {
                "total_sleep_logs": total_logs,
                "average_sleep_duration_hours": round(avg_sleep, 2),
                "schedule_adherence_days": adherence_days,
                "adherence_rate_percentage": round((adherence_days / len(scores)) * 100, 1) if len(scores) > 0 else 0.0
            },
            "data": [
                {
                    "date": log.date,
                    "sleep_duration_hours": log.sleep_duration,
                    "wake_up_time": log.wake_up_time,
                    "target_wake_up_time": log.target_wake_up_time
                } for log in behavior_logs
            ]
        }
        
    return {"error": "Invalid report type"}

def export_report_excel(report_type: str, data: Dict[str, Any]) -> io.BytesIO:
    """Generates a beautifully formatted Excel workbook binary."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report Details"
    
    # Enable gridlines
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    kpi_title_font = Font(name="Calibri", size=10, bold=True, color="595959")
    kpi_value_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    data_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)
    
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    kpi_fill = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    alt_row_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    
    border_thin = Side(border_style="thin", color="D9D9D9")
    border_double = Side(border_style="double", color="1F497D")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    bottom_total_border = Border(top=border_thin, bottom=border_double)
    
    # 1. Report Title
    ws["A1"] = data["report_name"]
    ws["A1"].font = title_font
    ws["A2"] = f"Period: {data['start_date']} to {data['end_date']} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True)
    
    # 2. Add KPI Summary Cards
    ws["A4"] = "SUMMARY METRICS"
    ws["A4"].font = Font(name="Calibri", size=12, bold=True, color="1F497D")
    
    summary_dict = data.get("summary", {})
    col_idx = 1
    for label, val in summary_dict.items():
        if isinstance(val, dict): # Skip breakdowns like category breakdown for Excel KPI cards
            continue
        
        # Write to columns
        lbl_cell = ws.cell(row=5, column=col_idx)
        val_cell = ws.cell(row=6, column=col_idx)
        
        clean_label = label.replace("_", " ").title()
        lbl_cell.value = clean_label
        lbl_cell.font = kpi_title_font
        lbl_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        lbl_cell.fill = kpi_fill
        lbl_cell.border = cell_border
        
        val_cell.value = val
        val_cell.font = kpi_value_font
        val_cell.alignment = Alignment(horizontal="center", vertical="center")
        val_cell.fill = kpi_fill
        val_cell.border = cell_border
        
        col_idx += 1
        
    # 3. Write Data Table
    table_start_row = 9
    ws.cell(row=table_start_row - 1, column=1).value = "DETAILED RECORD DATA"
    ws.cell(row=table_start_row - 1, column=1).font = Font(name="Calibri", size=12, bold=True, color="1F497D")
    
    list_data = data.get("data", [])
    if list_data:
        # Determine Headers from first item
        headers = [h.replace("_", " ").title() for h in list_data[0].keys()]
        
        # Write Headers
        for c_idx, h_name in enumerate(headers, 1):
            cell = ws.cell(row=table_start_row, column=c_idx)
            cell.value = h_name
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border
            
        # Write Rows
        current_row = table_start_row + 1
        for item in list_data:
            for c_idx, val in enumerate(item.values(), 1):
                cell = ws.cell(row=current_row, column=c_idx)
                
                # Format Boolean Values
                if isinstance(val, bool):
                    cell.value = "Yes" if val else "No"
                else:
                    cell.value = val
                    
                cell.font = data_font
                cell.border = cell_border
                
                # Alternate Row Coloring
                if current_row % 2 == 0:
                    cell.fill = alt_row_fill
                
                # Right align numbers
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(horizontal="left")
                    
            current_row += 1
            
        # Add averages at bottom if it is a habit report
        if report_type == "habit":
            ws.cell(row=current_row, column=1).value = "Average"
            ws.cell(row=current_row, column=1).font = bold_font
            ws.cell(row=current_row, column=1).border = bottom_total_border
            
            for c_idx in range(2, len(headers) + 1):
                col_letter = get_column_letter(c_idx)
                avg_cell = ws.cell(row=current_row, column=c_idx)
                avg_cell.value = f"=AVERAGE({col_letter}{table_start_row+1}:{col_letter}{current_row-1})"
                avg_cell.font = bold_font
                avg_cell.border = bottom_total_border
                avg_cell.alignment = Alignment(horizontal="right")
                avg_cell.number_format = "0.00"
                
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Ignore title and period rows for width calculation
            if cell.row < 4:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # Write to buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_report_pdf(report_type: str, data: Dict[str, Any], user_name: str) -> io.BytesIO:
    """Generates a beautifully styled PDF report using reportlab."""
    buffer = io.BytesIO()
    
    # 0.75 in margins
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    primary_color = colors.HexColor("#1F497D")
    text_color = colors.HexColor("#333333")
    light_grey = colors.HexColor("#F5F5F5")
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=primary_color,
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=10,
        leading=12,
        textColor=colors.HexColor("#555555"),
        spaceAfter=15
    )
    
    section_style = ParagraphStyle(
        'SectionHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=primary_color,
        spaceBefore=15,
        spaceAfter=8
    )
    
    body_style = ParagraphStyle(
        'BodyText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=text_color
    )
    
    tbl_hdr_style = ParagraphStyle(
        'TableHead',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.white,
        alignment=1 # Center
    )
    
    tbl_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=text_color,
        alignment=0 # Left
    )
    
    tbl_cell_center = ParagraphStyle(
        'TableCellCenter',
        parent=tbl_cell_style,
        alignment=1 # Center
    )
    
    story = []
    
    # Header block
    story.append(Paragraph(data["report_name"], title_style))
    meta_text = f"User: <b>{user_name}</b> | Period: {data['start_date']} to {data['end_date']} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    story.append(Paragraph(meta_text, subtitle_style))
    story.append(Spacer(1, 10))
    
    # KPI metrics block
    story.append(Paragraph("Executive Performance Summary", section_style))
    summary_dict = data.get("summary", {})
    
    kpi_data = []
    kpi_row_labels = []
    kpi_row_vals = []
    
    count = 0
    for label, val in summary_dict.items():
        if isinstance(val, dict): # Skip breakdowns in top level cards
            continue
        clean_label = label.replace("_", " ").title()
        kpi_row_labels.append(Paragraph(f"<b>{clean_label}</b>", ParagraphStyle('KPIlbl', parent=body_style, alignment=1, fontSize=8)))
        kpi_row_vals.append(Paragraph(f"<font size=12 color='#1F497D'><b>{val}</b></font>", ParagraphStyle('KPIval', parent=body_style, alignment=1)))
        count += 1
        
        # Chunk into rows of 3 columns
        if count == 3:
            kpi_data.append(kpi_row_labels)
            kpi_data.append(kpi_row_vals)
            kpi_row_labels = []
            kpi_row_vals = []
            count = 0
            
    if count > 0:
        # Pad remaining columns to have exactly 3 columns
        while len(kpi_row_labels) < 3:
            kpi_row_labels.append("")
            kpi_row_vals.append("")
        kpi_data.append(kpi_row_labels)
        kpi_data.append(kpi_row_vals)
        
    kpi_table = Table(kpi_data, colWidths=[2.2 * inch, 2.2 * inch, 2.2 * inch])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), light_grey),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('LINEBELOW', (0,0), (-1,0), 0.5, colors.white),
        ('BOX', (0,0), (-1,-1), 1, primary_color),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.white)
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 15))
    
    # If Category breakdown is present (for challenges), show it
    cat_breakdown = summary_dict.get("category_breakdown", {})
    if cat_breakdown:
        story.append(Paragraph("Category Performance Breakdown", section_style))
        cat_data = [[Paragraph("<b>Challenge Category</b>", tbl_cell_style), Paragraph("<b>Tasks Solved</b>", tbl_cell_center)]]
        for cat, val in cat_breakdown.items():
            cat_data.append([Paragraph(cat, tbl_cell_style), Paragraph(str(val), tbl_cell_center)])
        cat_table = Table(cat_data, colWidths=[4.0 * inch, 2.6 * inch])
        cat_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), primary_color),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, light_grey]),
            ('BOX', (0,0), (-1,-1), 1, primary_color),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#D9D9D9")),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
            ('TOPPADDING', (0,0), (-1,-1), 6)
        ]))
        # Update text color of headers to white
        for col_i in range(2):
            cat_data[0][col_i].style = tbl_hdr_style
        story.append(cat_table)
        story.append(Spacer(1, 15))

    # Detailed Logs Section
    story.append(Paragraph("Detailed Performance Log Records", section_style))
    list_data = data.get("data", [])
    
    if list_data:
        headers = list(list_data[0].keys())
        
        # Format Headers
        pdf_headers = [Paragraph(f"<b>{h.replace('_', ' ').title()}</b>", tbl_hdr_style) for h in headers]
        pdf_rows = [pdf_headers]
        
        for item in list_data:
            row = []
            for k, val in item.items():
                if isinstance(val, bool):
                    txt = "Yes" if val else "No"
                elif isinstance(val, float):
                    txt = f"{val:.2f}"
                elif val is None:
                    txt = "-"
                else:
                    txt = str(val)
                row.append(Paragraph(txt, tbl_cell_center if k in ["date", "wake_up_time", "target_wake_up_time"] else tbl_cell_style))
            pdf_rows.append(row)
            
        # Dynamically set column widths
        col_count = len(headers)
        avail_width = 6.6 * inch
        width_per_col = avail_width / col_count
        
        # Specialized widths based on report type to ensure fit
        col_widths = [width_per_col] * col_count
        if report_type == "habit":
            col_widths = [1.1*inch, 1.1*inch, 1.1*inch, 1.1*inch, 1.1*inch, 1.1*inch]
        elif report_type == "wake-up":
            col_widths = [0.9*inch, 0.9*inch, 0.9*inch, 0.9*inch, 0.9*inch, 1.1*inch, 1.0*inch]
            
        log_table = Table(pdf_rows, colWidths=col_widths, repeatRows=1)
        log_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), primary_color),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, light_grey]),
            ('BOX', (0,0), (-1,-1), 1, primary_color),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#D9D9D9")),
            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
            ('TOPPADDING', (0,0), (-1,-1), 5)
        ]))
        story.append(log_table)
    else:
        story.append(Paragraph("No record entries found for the selected period.", body_style))
        
    doc.build(story)
    buffer.seek(0)
    return buffer
